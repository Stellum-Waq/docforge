"""OCR 动作的真实集成测试。

与其它测试不同，这里**真的调用本地 OCR 引擎**（RapidOCR + ONNX），
因此能验证"图片 → 文字 / 表格"这条链路确实可用，而不是只验证接口形状。

引擎是模块级单例，模型只在首次调用时加载（约 0.5s），后续复用。
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from docforge.actions import ActionContext, get_action
from docforge.jobs import JobRequest
from docforge.jobs import manager as manager_module
from docforge.ocr import OcrOptions, engine_status, recognize
from docforge.ocr.table import blocks_to_grid

FONT_PATH = "C:/Windows/Fonts/msyh.ttc"
pytestmark = pytest.mark.skipif(
    not Path(FONT_PATH).is_file(), reason="测试需要微软雅黑字体"
)


# --------------------------------------------------------------------------- #
# 素材                                                                          #
# --------------------------------------------------------------------------- #

@pytest.fixture(scope="module")
def document_image(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """一张包含标题与段落的文档图。"""
    directory = tmp_path_factory.mktemp("ocr")
    image = Image.new("RGB", (1000, 460), "white")
    draw = ImageDraw.Draw(image)

    title = ImageFont.truetype(FONT_PATH, 40)
    body = ImageFont.truetype(FONT_PATH, 26)

    draw.text((50, 40), "文枢 DocForge 文字识别测试", font=title, fill="black")
    draw.text((50, 130), "第一行：办公文件转换工具", font=body, fill="black")
    draw.text((50, 180), "Second line: mixed English 12345", font=body, fill="black")
    draw.text((50, 230), "金额：¥1,234.56", font=body, fill="black")
    draw.text((50, 280), "日期：2026-09-11", font=body, fill="black")

    path = directory / "document.png"
    image.save(path)
    return path


@pytest.fixture(scope="module")
def table_image(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """一张带网格线的三列表格图。"""
    directory = tmp_path_factory.mktemp("ocr")
    image = Image.new("RGB", (900, 320), "white")
    draw = ImageDraw.Draw(image)

    font = ImageFont.truetype(FONT_PATH, 28)
    headers = ["姓名", "部门", "金额"]
    rows = [["张三", "技术部", "1200"], ["李四", "市场部", "980"]]

    column_x = [30, 300, 570, 840]
    row_y = [40, 130, 220, 310]

    for y in row_y:
        draw.line([(column_x[0], y), (column_x[-1], y)], fill="black", width=3)
    for x in column_x:
        draw.line([(x, row_y[0]), (x, row_y[-1])], fill="black", width=3)

    for col, text in enumerate(headers):
        draw.text((column_x[col] + 30, row_y[0] + 30), text, font=font, fill="black")
    for r, values in enumerate(rows, start=1):
        for col, text in enumerate(values):
            draw.text((column_x[col] + 30, row_y[r] + 30), text, font=font, fill="black")

    path = directory / "table.png"
    image.save(path)
    return path


# --------------------------------------------------------------------------- #
# 引擎层                                                                        #
# --------------------------------------------------------------------------- #

def test_local_engine_is_available() -> None:
    status = {item["id"]: item for item in engine_status()}
    assert status["ocr.rapidocr"]["available"] is True


def test_local_ocr_recognizes_chinese_and_english(document_image: Path) -> None:
    result = recognize(document_image, OcrOptions(mode="text"), preference="local")

    assert result.engine == "ocr.rapidocr"
    assert len(result.blocks) >= 4
    assert result.average_confidence > 0.7

    text = result.text.replace(" ", "")
    # 逐项断言关键内容，任何一项丢了都说明识别质量出问题
    assert "文枢" in text
    assert "DocForge" in text
    assert "办公文件转换工具" in text
    assert "Secondline" in text.replace(" ", "")
    assert "12345" in text
    assert "2026-09-11" in text


def test_local_ocr_blocks_have_reading_order(document_image: Path) -> None:
    """结果必须按从上到下的阅读顺序返回，否则拼接出来的文本是乱的。"""
    result = recognize(document_image, OcrOptions(mode="text"), preference="local")
    tops = [block.box[1] for block in result.blocks]
    assert tops == sorted(tops)


def test_local_ocr_detects_table_columns(table_image: Path) -> None:
    result = recognize(table_image, OcrOptions(mode="text"), preference="local")
    grid = blocks_to_grid(result.blocks)

    assert grid.column_count == 3, f"应识别出 3 列，实际 {grid.column_count}：{grid.rows}"
    assert grid.row_count == 3

    flat = ["".join(row) for row in grid.rows]
    assert any("姓名" in row for row in flat)
    assert any("1200" in row for row in flat)


def test_cloud_engine_reports_unavailable_without_key() -> None:
    """没有配置密钥时，云端引擎必须给出可读原因而不是抛异常。"""
    from docforge.ocr.deepseek import get_engine

    available, reason = get_engine().availability()
    assert available is False
    assert reason and "API Key" in reason


def test_router_falls_back_when_cloud_unavailable(document_image: Path) -> None:
    """即使要求走云端，云端不可用时也应自动降级到本地而不是直接失败。"""
    result = recognize(document_image, OcrOptions(mode="text"), preference="cloud")
    assert result.engine == "ocr.rapidocr"
    assert any("云端" in warning for warning in result.warnings)


# --------------------------------------------------------------------------- #
# 动作层                                                                        #
# --------------------------------------------------------------------------- #

def _run_action(action_id: str, source: Path, out_dir: Path, params: dict, name: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / name
    ctx = ActionContext(
        job_id="test",
        task_id="t",
        file_path=str(source),
        output_path=str(target),
        params=params,
    )
    result = get_action(action_id).handler(ctx)
    assert result.output_path
    return Path(result.output_path)


def test_ocr_action_writes_text(document_image: Path, tmp_path: Path) -> None:
    out = _run_action(
        "image.ocr",
        document_image,
        tmp_path / "out",
        {"engine": "local", "mode": "text", "output_format": "txt"},
        "document.txt",
    )

    assert out.suffix == ".txt"
    content = out.read_text(encoding="utf-8")
    assert "文枢" in content
    assert "2026-09-11" in content


def test_ocr_action_writes_markdown(document_image: Path, tmp_path: Path) -> None:
    out = _run_action(
        "image.ocr",
        document_image,
        tmp_path / "out",
        {"engine": "local", "mode": "layout", "output_format": "md"},
        "document.md",
    )
    assert out.suffix == ".md"
    assert len(out.read_text(encoding="utf-8").strip()) > 10


def test_ocr_action_writes_docx(document_image: Path, tmp_path: Path) -> None:
    out = _run_action(
        "image.ocr",
        document_image,
        tmp_path / "out",
        {"engine": "local", "mode": "text", "output_format": "docx"},
        "document.docx",
    )
    assert out.suffix == ".docx"
    assert out.stat().st_size > 0

    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "文枢" in text


def test_ocr_action_writes_json(document_image: Path, tmp_path: Path) -> None:
    import json

    out = _run_action(
        "image.ocr",
        document_image,
        tmp_path / "out",
        {"engine": "local", "mode": "text", "output_format": "json", "include_boxes": True},
        "document.json",
    )
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["engine"] == "ocr.rapidocr"
    assert payload["lineCount"] >= 4
    assert payload["blocks"][0]["box"]


def test_image_to_excel_action_builds_workbook(table_image: Path, tmp_path: Path) -> None:
    out = _run_action(
        "image.to_excel",
        table_image,
        tmp_path / "out",
        {"engine": "local", "header_row": "auto", "sheet_mode": "single"},
        "table.xlsx",
    )

    assert out.suffix == ".xlsx"
    workbook = load_workbook(out)
    sheet = workbook.active

    # 表头 + 2 行数据
    assert sheet.max_row == 3, f"应有 3 行，实际 {sheet.max_row}"
    assert sheet.max_column == 3, f"应有 3 列，实际 {sheet.max_column}"

    values = [[sheet.cell(row=r, column=c).value for c in range(1, 4)] for r in range(1, 4)]
    flat = ["".join(str(v) for v in row if v is not None) for row in values]

    assert "姓名" in flat[0]
    assert any("张三" in row for row in flat)
    assert any("1200" in row for row in flat)


def test_image_to_excel_action_can_include_text_sheet(table_image: Path, tmp_path: Path) -> None:
    out = _run_action(
        "image.to_excel",
        table_image,
        tmp_path / "out",
        {"engine": "local", "include_text": True, "include_source_sheet": True},
        "table_full.xlsx",
    )
    workbook = load_workbook(out)
    assert "原文" in workbook.sheetnames
    assert "识别信息" in workbook.sheetnames


def test_image_to_excel_never_loses_lines(document_image: Path, tmp_path: Path) -> None:
    """把无表格结构的文档图转 Excel 时，**一行都不能丢**。

    这里守的是一个真实踩过的坑：早先的实现先"猜哪一行是表头"，然后把表头
    **之前**的行全部丢弃。文档正文的第一行经常被误判成表头，结果用户拿到
    的 Excel 凭空少了几行，而且没有任何提示 —— 静默数据丢失是最难排查的一类 bug。
    """
    out = _run_action(
        "image.to_excel",
        document_image,
        tmp_path / "out",
        {"engine": "local", "header_row": "auto"},
        "document.xlsx",
    )

    sheet = load_workbook(out).active
    cells = [
        str(sheet.cell(row=r, column=c).value)
        for r in range(1, sheet.max_row + 1)
        for c in range(1, sheet.max_column + 1)
        if sheet.cell(row=r, column=c).value not in (None, "")
    ]
    joined = "".join(cells)

    # document.png 里有 5 行文字，全部都要出现在 Excel 中
    assert "文枢" in joined, f"标题行丢失：{cells}"
    assert "办公文件转换工具" in joined, f"正文丢失：{cells}"
    assert "12345" in joined or "Secondline" in joined.replace(" ", ""), f"英文行丢失：{cells}"
    assert "1,234.56" in joined, f"金额行丢失：{cells}"
    assert "2026-09-11" in joined, f"日期行丢失：{cells}"


def test_image_to_excel_header_none_keeps_all_rows(table_image: Path, tmp_path: Path) -> None:
    """header_row=none 时所有行都当数据，行数不应减少。"""
    out = _run_action(
        "image.to_excel",
        table_image,
        tmp_path / "out",
        {"engine": "local", "header_row": "none"},
        "no_header.xlsx",
    )
    sheet = load_workbook(out).active
    assert sheet.max_row == 3
    assert sheet.max_column == 3


# --------------------------------------------------------------------------- #
# 队列层                                                                        #
# --------------------------------------------------------------------------- #

@pytest.fixture()
async def manager():
    instance = manager_module.JobManager()
    manager_module._manager = instance
    await instance.start()
    try:
        yield instance
    finally:
        await instance.shutdown()
        manager_module._manager = None


async def _wait(job_id: str, timeout: float = 90.0):
    deadline = asyncio.get_running_loop().time() + timeout
    while asyncio.get_running_loop().time() < deadline:
        job = manager_module.get_manager().get(job_id)
        if job and job.status in {"succeeded", "failed", "cancelled"}:
            return job
        await asyncio.sleep(0.1)
    raise AssertionError("任务未在超时内完成")


async def test_ocr_through_job_queue(manager, document_image: Path, tmp_path: Path) -> None:
    """走完整任务队列：批量 OCR 多张图，验证进度与产物。"""
    files = [str(document_image), str(document_image)]

    job = manager.create_job(
        JobRequest(
            action="image.ocr",
            files=files,
            output_dir=str(tmp_path / "queue-out"),
            params={"engine": "local", "mode": "text", "output_format": "txt"},
            concurrency=2,
        )
    )
    manager.start_job(job)
    settled = await _wait(job.id)

    assert settled.status == "succeeded"
    assert settled.completed == 2
    assert settled.progress == 100.0

    produced = sorted((tmp_path / "queue-out").glob("*.txt"))
    assert len(produced) == 2
    for path in produced:
        assert "文枢" in path.read_text(encoding="utf-8")

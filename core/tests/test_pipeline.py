"""多步骤流水线的测试。

流水线把一串动作串起来跑，风险集中在三处，因此测试也围绕这三处：

1. **类型链**：第 3 步才发现"上一步产出的是 pdf、这一步要 docx"，等于白跑前面
   两步（可能还包括几分钟的 OCR）。所以整条链必须在开工前就校验完。
2. **输出命名**：队列按动作声明的输出扩展名提前定好文件名。流水线的输出类型
   取决于最后一步，若沿用源文件扩展名，就会产出"内容是 PDF、后缀是 .docx"
   的文件 —— 用户双击打不开，还很难想明白问题在哪。
3. **中间产物**：中间文件必须写在临时目录并在结束后清干净。否则用户会在输出
   目录里看到一堆分不清是半成品还是结果的文件。
"""

from __future__ import annotations

import asyncio
import tempfile
from pathlib import Path

import pymupdf
import pytest
from PIL import Image

from docforge.actions import (
    ActionError,
    PipelineStep,
    describe_chain,
    get_action,
    parse_steps,
    validate_chain,
)
from docforge.actions.pipeline import STEPS_KEY, _pipeline_output_ext, run_chain
from docforge.jobs import JobRequest, get_manager
from docforge.jobs import manager as manager_module


# --------------------------------------------------------------------------- #
# 夹具与工具                                                                    #
# --------------------------------------------------------------------------- #

@pytest.fixture()
async def manager():
    instance = manager_module.JobManager()
    manager_module._manager = instance
    await instance.start()
    try:
        yield instance
    finally:
        await instance.shutdown()
        manager_module._manager = None


def _png(path: Path, size: tuple[int, int] = (240, 160)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (28, 44, 78)).save(path)
    return path


def _pdf(path: Path, pages: int = 3) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page(width=595, height=842)
    doc.save(str(path))
    doc.close()
    return path


def _docx(path: Path) -> Path:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("流水线测试", 0)
    doc.add_paragraph("用于验证 Word → PDF → 水印 的链式转换。")
    doc.save(str(path))
    return path


async def _wait(job_id: str, timeout: float = 120.0):
    mgr = get_manager()
    deadline = asyncio.get_running_loop().time() + timeout
    while asyncio.get_running_loop().time() < deadline:
        job = mgr.get(job_id)
        if job and job.status in {"succeeded", "failed", "cancelled"}:
            return job
        await asyncio.sleep(0.05)
    raise AssertionError(f"任务 {job_id} 在 {timeout}s 内未结束")


def _temp_pipeline_dirs() -> set[str]:
    return {p.name for p in Path(tempfile.gettempdir()).glob("docforge-pipeline-*")}


# --------------------------------------------------------------------------- #
# 步骤解析                                                                      #
# --------------------------------------------------------------------------- #

def test_parse_steps_accepts_plain_dicts() -> None:
    steps = parse_steps([{"action": "pdf.compress", "params": {"subset_fonts": False}}])
    assert steps == [PipelineStep("pdf.compress", {"subset_fonts": False})]


def test_parse_steps_defaults_params_to_empty() -> None:
    assert parse_steps([{"action": "pdf.compress"}]) == [PipelineStep("pdf.compress", {})]


@pytest.mark.parametrize(
    "raw",
    [
        None,
        [],
        "pdf.compress",
        [{"params": {}}],                       # 没有 action
        [{"action": "  "}],                      # action 是空白
        [{"action": "pdf.compress", "params": "x"}],  # params 不是对象
        [42],
    ],
)
def test_parse_steps_rejects_malformed_input(raw: object) -> None:
    with pytest.raises(ActionError):
        parse_steps(raw)


def test_parse_steps_enforces_step_limit() -> None:
    from docforge.actions.pipeline import MAX_STEPS

    steps = [{"action": "pdf.compress"} for _ in range(MAX_STEPS + 1)]
    with pytest.raises(ActionError, match="最多"):
        parse_steps(steps)

    # 刚好到上限应当通过
    assert len(parse_steps(steps[:MAX_STEPS])) == MAX_STEPS


# --------------------------------------------------------------------------- #
# 类型链校验                                                                    #
# --------------------------------------------------------------------------- #

def test_validate_chain_accepts_a_coherent_chain(tmp_path: Path) -> None:
    steps = parse_steps([
        {"action": "pdf.searchable", "params": {"dpi": 100}},
        {"action": "pdf.watermark", "params": {"text": "机密"}},
        {"action": "pdf.compress"},
    ])
    specs = validate_chain(str(_pdf(tmp_path / "a.pdf")), steps)
    assert [s.id for s in specs] == ["pdf.searchable", "pdf.watermark", "pdf.compress"]


def test_validate_chain_catches_mismatch_on_first_step(tmp_path: Path) -> None:
    steps = parse_steps([{"action": "pdf.watermark"}])
    with pytest.raises(ActionError, match="源文件是 .png"):
        validate_chain(str(_png(tmp_path / "a.png")), steps)


def test_validate_chain_catches_mismatch_between_steps(tmp_path: Path) -> None:
    """第 1 步把 pdf 变成 png，第 2 步要的是 pdf —— 必须报出是第几步。"""
    steps = parse_steps([
        {"action": "pdf.to_images", "params": {"image_format": "png"}},
        {"action": "pdf.watermark"},
    ])
    with pytest.raises(ActionError, match="第 2 步"):
        validate_chain(str(_pdf(tmp_path / "a.pdf")), steps)


def test_validate_chain_rejects_aggregate_step(tmp_path: Path) -> None:
    steps = parse_steps([{"action": "pdf.merge"}])
    with pytest.raises(ActionError, match="聚合动作"):
        validate_chain(str(_pdf(tmp_path / "a.pdf")), steps)


def test_validate_chain_rejects_nested_pipeline(tmp_path: Path) -> None:
    steps = parse_steps([{"action": "pipeline"}])
    with pytest.raises(ActionError, match="不允许嵌套"):
        validate_chain(str(_pdf(tmp_path / "a.pdf")), steps)


def test_validate_chain_reports_unknown_action(tmp_path: Path) -> None:
    steps = parse_steps([{"action": "no.such.action"}])
    with pytest.raises(ActionError, match="不存在"):
        validate_chain(str(_pdf(tmp_path / "a.pdf")), steps)


def test_validate_chain_keeps_extension_when_output_ext_is_none(tmp_path: Path) -> None:
    """图片水印的 output_ext 是 None（同源），链上类型不应被改写。"""
    steps = parse_steps([
        {"action": "image.watermark", "params": {"text": "x"}},
        {"action": "image.watermark", "params": {"text": "y"}},
    ])
    specs = validate_chain(str(_png(tmp_path / "a.png")), steps)
    assert all(s.output_ext is None for s in specs)


def test_describe_chain_shows_the_type_flow(tmp_path: Path) -> None:
    steps = parse_steps([{"action": "doc.to_pdf"}, {"action": "pdf.watermark"}])
    lines = describe_chain(str(_docx(tmp_path / "a.docx")), steps)
    assert lines == ["1. 文档转 PDF  .docx → .pdf", "2. PDF 加水印  .pdf → .pdf"]


# --------------------------------------------------------------------------- #
# 动态输出扩展名                                                                #
# --------------------------------------------------------------------------- #

def test_pipeline_output_ext_follows_the_last_step() -> None:
    assert _pipeline_output_ext({STEPS_KEY: [{"action": "doc.to_pdf"}]}) == "pdf"
    assert _pipeline_output_ext(
        {STEPS_KEY: [{"action": "doc.to_pdf"}, {"action": "pdf.info"}]}
    ) == "json"
    # 最后一步同源（图片水印）→ 不声明扩展名
    assert _pipeline_output_ext({STEPS_KEY: [{"action": "image.watermark"}]}) is None


def test_pipeline_output_ext_survives_garbage() -> None:
    """钩子只是为了让输出名好看，拿到坏数据不该抛异常把任务带崩。"""
    assert _pipeline_output_ext({}) is None
    assert _pipeline_output_ext({STEPS_KEY: "不是列表"}) is None
    assert _pipeline_output_ext({STEPS_KEY: [{"action": "no.such.action"}]}) is None


def test_action_spec_resolves_dynamic_ext(tmp_path: Path) -> None:
    spec = get_action("pipeline")
    assert spec.resolve_output_ext({STEPS_KEY: [{"action": "doc.to_pdf"}]}) == "pdf"
    # 解析不出来时退回静态值（None = 同源）
    assert spec.resolve_output_ext({}) is None


# --------------------------------------------------------------------------- #
# 执行                                                                          #
# --------------------------------------------------------------------------- #

def test_run_chain_single_step_writes_to_output_dir(tmp_path: Path) -> None:
    source = _png(tmp_path / "in" / "photo.png")
    out_dir = tmp_path / "out"

    result = run_chain(
        str(source),
        parse_steps([{"action": "image.watermark", "params": {"text": "机密", "opacity": 0.7}}]),
        str(out_dir),
    )

    produced = Path(result.output_path)
    assert produced.parent == out_dir
    assert produced.suffix == ".png"
    assert produced.read_bytes() != source.read_bytes()


def test_run_chain_three_steps_produce_the_final_result(tmp_path: Path) -> None:
    source = _pdf(tmp_path / "in" / "doc.pdf", pages=3)
    out_dir = tmp_path / "out"

    result = run_chain(
        str(source),
        parse_steps([
            {"action": "pdf.watermark", "params": {"text": "内部", "opacity": 0.3}},
            {"action": "pdf.rotate", "params": {"angle": "90"}},
            {"action": "pdf.extract_pages", "params": {"pages": "1-2"}},
        ]),
        str(out_dir),
    )

    produced = Path(result.output_path)
    assert produced.parent == out_dir
    # 输出目录里只有最终产物，中间文件不该出现在这里
    assert [p.name for p in out_dir.iterdir()] == [produced.name]

    doc = pymupdf.open(str(produced))
    try:
        assert doc.page_count == 2, "extract_pages 是最后一步，页数应为 2"
        assert doc[0].rotation == 90
        assert "内部" in doc[0].get_text()
    finally:
        doc.close()


def test_run_chain_uses_final_target_verbatim(tmp_path: Path) -> None:
    """队列已经按重名策略解析过输出名，流水线必须原样使用，不能再解析一次。"""
    source = _png(tmp_path / "in" / "photo.png")
    target = tmp_path / "out" / "自定义名字.png"

    result = run_chain(
        str(source),
        parse_steps([{"action": "image.watermark", "params": {"text": "x"}}]),
        str(tmp_path / "out"),
        final_target=str(target),
    )

    assert Path(result.output_path) == target
    assert target.exists()


def test_run_chain_cleans_up_intermediate_files(tmp_path: Path) -> None:
    before = _temp_pipeline_dirs()

    run_chain(
        str(_pdf(tmp_path / "in" / "doc.pdf")),
        parse_steps([{"action": "pdf.compress"}, {"action": "pdf.rotate", "params": {"angle": "90"}}]),
        str(tmp_path / "out"),
    )

    assert _temp_pipeline_dirs() == before, "临时目录必须被清理，否则会越积越多"


def test_run_chain_cleans_up_on_failure(tmp_path: Path) -> None:
    """失败路径同样要清理 —— 报错时留下的临时目录最容易被忽略。"""
    before = _temp_pipeline_dirs()

    # 第 2 步必然失败：一个 240×160 的图片被当成 PDF 处理
    source = _png(tmp_path / "in" / "photo.png")
    with pytest.raises(ActionError):
        run_chain(
            str(source),
            parse_steps([
                {"action": "image.watermark", "params": {"text": "x"}},
                {"action": "pdf.info"},
            ]),
            str(tmp_path / "out"),
        )

    assert _temp_pipeline_dirs() == before


def test_run_chain_error_names_the_failing_step(tmp_path: Path) -> None:
    source = _pdf(tmp_path / "in" / "doc.pdf")
    with pytest.raises(ActionError, match="第 2 步"):
        run_chain(
            str(source),
            parse_steps([
                {"action": "pdf.compress"},
                # pages 越界：提取不存在的页会让动作报错
                {"action": "pdf.extract_pages", "params": {"pages": "99-100"}},
            ]),
            str(tmp_path / "out"),
        )


def test_run_chain_message_lists_the_steps(tmp_path: Path) -> None:
    result = run_chain(
        str(_pdf(tmp_path / "in" / "doc.pdf")),
        parse_steps([{"action": "pdf.compress"}, {"action": "pdf.compress"}]),
        str(tmp_path / "out"),
    )
    assert result.message == "PDF 压缩 → PDF 压缩"


def test_run_chain_reports_progress_per_step(tmp_path: Path) -> None:
    seen: list[float] = []
    run_chain(
        str(_pdf(tmp_path / "in" / "doc.pdf")),
        parse_steps([{"action": "pdf.compress"}, {"action": "pdf.compress"}]),
        str(tmp_path / "out"),
        report=lambda progress, _message: seen.append(progress),
    )
    assert seen, "每一步都应上报进度，否则界面一直转圈"
    assert seen[-1] == 100.0
    assert max(seen) <= 100.0


# --------------------------------------------------------------------------- #
# 经过任务队列（真实路径）                                                       #
# --------------------------------------------------------------------------- #

@pytest.mark.asyncio
async def test_pipeline_job_through_manager(manager, tmp_path: Path) -> None:
    """走真实队列：多文件、并发、逐文件归因、输出扩展名。"""
    sources = [_pdf(tmp_path / "in" / f"doc{i}.pdf", pages=2) for i in range(3)]
    out_dir = tmp_path / "out"

    request = JobRequest(
        action="pipeline",
        files=[str(p) for p in sources],
        output_dir=str(out_dir),
        params={
            STEPS_KEY: [
                {"action": "pdf.watermark", "params": {"text": "机密", "opacity": 0.3}},
                {"action": "pdf.rotate", "params": {"angle": "90"}},
            ]
        },
        concurrency=2,
    )
    job = manager.create_job(request)
    manager.start_job(job)
    finished = await _wait(job.id)

    assert finished.status == "succeeded"
    assert finished.failed == 0
    assert finished.completed == 3

    produced = sorted(out_dir.glob("*.pdf"))
    assert len(produced) == 3, "三个源文件应产出三个结果，不能互相覆盖"
    for path in produced:
        doc = pymupdf.open(str(path))
        try:
            assert doc[0].rotation == 90
            assert "机密" in doc[0].get_text()
        finally:
            doc.close()


@pytest.mark.asyncio
async def test_pipeline_job_output_extension_follows_last_step(manager, tmp_path: Path) -> None:
    """以「文档转 PDF」收尾时，产物必须是 .pdf。

    这是最容易出错的一点：队列提前用动作的 output_ext 定文件名，
    而流水线的 output_ext 是动态的。写错会让用户拿到一个内容是 PDF、
    后缀却是 .docx 的文件。
    """
    source = _docx(tmp_path / "in" / "report.docx")
    out_dir = tmp_path / "out"

    request = JobRequest(
        action="pipeline",
        files=[str(source)],
        output_dir=str(out_dir),
        params={STEPS_KEY: [{"action": "doc.to_pdf"}]},
        concurrency=1,
    )
    job = manager.create_job(request)
    manager.start_job(job)
    finished = await _wait(job.id)

    assert finished.status == "succeeded", finished.to_dict()

    produced = list(out_dir.iterdir())
    assert len(produced) == 1
    assert produced[0].suffix == ".pdf", f"输出扩展名应为 .pdf，实际是 {produced[0].name}"
    # 内容也得真的是 PDF
    doc = pymupdf.open(str(produced[0]))
    try:
        assert doc.page_count >= 1
    finally:
        doc.close()


@pytest.mark.asyncio
async def test_pipeline_job_failure_is_attributed_per_file(manager, tmp_path: Path) -> None:
    """一个文件失败不能拖垮整批，而且要说明是第几步出的问题。"""
    good = _pdf(tmp_path / "in" / "good.pdf")
    bad = tmp_path / "in" / "bad.pdf"
    bad.write_bytes(b"%PDF-1.4 broken")     # 损坏的 PDF：第 1 步就该失败

    out_dir = tmp_path / "out"
    request = JobRequest(
        action="pipeline",
        files=[str(good), str(bad)],
        output_dir=str(out_dir),
        params={STEPS_KEY: [{"action": "pdf.compress"}]},
        concurrency=1,
    )
    job = manager.create_job(request)
    manager.start_job(job)
    finished = await _wait(job.id)

    assert finished.completed == 1
    assert finished.failed == 1

    by_name = {t.file_name: t for t in finished.tasks}
    assert by_name["good.pdf"].status == "succeeded"
    assert by_name["bad.pdf"].status == "failed"
    assert by_name["bad.pdf"].error

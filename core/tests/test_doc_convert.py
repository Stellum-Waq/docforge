"""文档转换动作的测试（Word ↔ PDF）。

这是**真实调用 Office COM** 的集成测试，因此比其它测试慢（每次要拉起 Word 进程）。
但值得：COM 自动化有三个经典坑，只有真跑一遍才能发现 ——

  1. COM 是 STA，对象只能在创建它的线程上用
  2. 弹出模态对话框会让调用永久挂起，必须靠看门狗超时强杀
  3. 用完不清理会在用户机器上留下无主的 WINWORD.EXE

因此下面的用例专门盯着第 3 点：转换完成后**不能有残留进程**。
"""

from __future__ import annotations

import sys
from pathlib import Path

import psutil
import pymupdf
import pytest

from docforge.actions import ActionContext, ActionError, get_action
from docforge.engines.com import shutdown_all

WINDOWS = sys.platform == "win32"
HAS_WORD = False
if WINDOWS:
    try:
        from docforge.engines.probe import _office_exe

        HAS_WORD = _office_exe("word") is not None
    except Exception:  # noqa: BLE001
        HAS_WORD = False

requires_word = pytest.mark.skipif(
    not (WINDOWS and HAS_WORD), reason="需要 Windows + 已安装 Microsoft Word"
)


# --------------------------------------------------------------------------- #
# 素材                                                                          #
# --------------------------------------------------------------------------- #

def _make_docx(path: Path, *, title: str = "季度工作报告", paragraphs: int = 4) -> Path:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    document.add_heading("一、总体情况", level=2)
    for index in range(paragraphs):
        document.add_paragraph(f"第 {index + 1} 段：本季度各项指标均按计划推进，办公文件转换工具已完成核心功能。")

    table = document.add_table(rows=3, cols=3)
    for col, header in enumerate(("项目", "计划", "完成")):
        table.cell(0, col).text = header
    rows = (("文档转换", "100", "100"), ("文字识别", "80", "85"))
    for row_index, values in enumerate(rows, start=1):
        for col, value in enumerate(values):
            table.cell(row_index, col).text = value

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def _run(action_id: str, source: Path, out_dir: Path, params: dict, name: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / name
    ctx = ActionContext(
        job_id="test",
        task_id="t",
        file_path=str(source),
        output_path=str(target),
        params=params,
    )
    result = get_action(action_id).handler(ctx)
    assert result.output_path
    return Path(result.output_path)


def _word_pids() -> set[int]:
    pids = set()
    for process in psutil.process_iter(["name"]):
        if (process.info.get("name") or "").lower() == "winword.exe":
            pids.add(process.pid)
    return pids


# --------------------------------------------------------------------------- #
# 注册与元数据                                                                  #
# --------------------------------------------------------------------------- #

def test_actions_are_registered() -> None:
    from docforge.actions import list_actions

    ids = {spec.id for spec in list_actions()}
    assert "doc.to_pdf" in ids
    assert "pdf.to_word" in ids


def test_to_pdf_accepts_office_formats() -> None:
    spec = get_action("doc.to_pdf")
    for ext in ("docx", "doc", "xlsx", "pptx", "rtf", "txt", "md"):
        assert ext in spec.accepts, f"应接受 .{ext}"


# --------------------------------------------------------------------------- #
# Word → PDF（COM 高保真路径）                                                  #
# --------------------------------------------------------------------------- #

@requires_word
def test_docx_to_pdf_via_com(tmp_path: Path) -> None:
    """COM 路径必须产出**带文字层**的 PDF —— 这是"高保真"的核心证据。"""
    source = _make_docx(tmp_path / "报告.docx")
    out = _run("doc.to_pdf", source, tmp_path / "out", {"engine": "com"}, "报告.pdf")

    assert out.is_file() and out.stat().st_size > 5000, "COM 产出的 PDF 不应过小"

    doc = pymupdf.open(str(out))
    try:
        text = "".join(page.get_text() for page in doc)
        assert "季度工作报告" in text, "标题应被还原"
        assert "总体情况" in text, "二级标题应被还原"
        assert "办公文件转换工具" in text, "正文应被还原"
        assert "文档转换" in text, "表格内容应被还原"
    finally:
        doc.close()


@requires_word
def test_com_generates_pdf_bookmarks(tmp_path: Path) -> None:
    """开启书签后，PDF 应带出目录结构（Word 的标题样式 → PDF 书签）。"""
    source = _make_docx(tmp_path / "报告.docx")
    out = _run(
        "doc.to_pdf", source, tmp_path / "out", {"engine": "com", "bookmarks": True}, "报告.pdf"
    )

    doc = pymupdf.open(str(out))
    try:
        toc = doc.get_toc()
        assert toc, "应生成书签目录"
        titles = [entry[1] for entry in toc]
        assert any("季度工作报告" in title for title in titles)
    finally:
        doc.close()


@requires_word
def test_no_orphan_word_process_after_conversion(tmp_path: Path) -> None:
    """转换结束后不能留下无主的 WINWORD.EXE。

    这是 COM 自动化最常见、也最惹人烦的后遗症：用户看不到窗口，
    进程却在后台占内存，还可能锁住文档导致下次转换失败。

    注意基线要在 ``shutdown_all()`` **之后**取 —— 上一条用例可能留下进程，
    拿它当基线会让断言失去意义。
    """
    shutdown_all()
    baseline = _word_pids()

    source = _make_docx(tmp_path / "报告.docx")
    _run("doc.to_pdf", source, tmp_path / "out", {"engine": "com"}, "报告.pdf")

    # 主动关停工作线程（模拟应用退出）
    shutdown_all()

    leaked = _word_pids() - baseline
    assert not leaked, f"残留了 Word 进程：{leaked}"


@requires_word
def test_repeated_conversions_reuse_the_same_worker(tmp_path: Path) -> None:
    """连续转换多个文件应复用同一个 Word 实例，而不是每次重启。

    重启 Word 每次要 1–3 秒，批量处理上百个文件时差别非常明显。
    """
    shutdown_all()
    baseline = _word_pids()
    sources = [_make_docx(tmp_path / f"doc{i}.docx", title=f"文档{i}") for i in range(3)]

    for index, source in enumerate(sources):
        out = _run(
            "doc.to_pdf", source, tmp_path / f"out{index}", {"engine": "com"}, f"doc{index}.pdf"
        )
        assert out.is_file()

    # 过程当中应当只有一个新增的 Word 进程（复用了同一个实例）
    during = _word_pids() - baseline
    assert len(during) <= 1, f"不应为每个文件重启 Word，当前新增 {len(during)} 个进程"

    shutdown_all()
    assert not (_word_pids() - baseline), "关停后不应有残留"


@requires_word
def test_failed_call_does_not_orphan_a_word_process(tmp_path: Path) -> None:
    """**回归测试**：一次失败的 COM 调用不能留下再也清理不掉的孤儿进程。

    这条用例来自一个真实缺陷。``_invalidate()`` 原本只把应用引用置空、
    不同步杀掉进程，于是：

      * 每次调用失败都会留下一个 WINWORD.EXE；
      * 更糟的是，下一次 ``_get_app()`` 会把"启动前已存在的进程"整体记为
        **外来进程**，那个孤儿从此不在 ``_owned_pids`` 里 —— 之后的
        ``shutdown_all()``（以及内核退出时的兜底清理）再也认领不到它。

    实测复现：连续 4 次失败 → 4 个残留 WINWORD.EXE，``shutdown_all()``
    之后依然全部存活。用户侧的表现是"用久了后台一堆 Word 进程"，
    而且关掉软件也清不掉。
    """
    from docforge.engines.com import ComAppWorker

    shutdown_all()
    baseline = _word_pids()

    worker = ComAppWorker(
        prog_id="Word.Application",
        process_name="WINWORD.EXE",
        label="Microsoft Word",
    )
    try:
        for _ in range(3):
            with pytest.raises(Exception):
                # 每次调用都抛异常，逼工作线程走 _invalidate() 分支
                worker.run(
                    lambda app: (_ for _ in ()).throw(RuntimeError("模拟转换失败")),
                    timeout=60,
                )
            leaked = _word_pids() - baseline
            assert not leaked, f"调用失败后残留了 Word 进程：{leaked}"
    finally:
        shutdown_all()

    assert not (_word_pids() - baseline), "关停后不应有残留"


# --------------------------------------------------------------------------- #
# 内置渲染（降级路径）                                                          #
# --------------------------------------------------------------------------- #

def test_builtin_engine_works_without_office(tmp_path: Path) -> None:
    """内置渲染是最后一道防线：没有 Office 也要能用，哪怕是低保真。"""
    source = _make_docx(tmp_path / "报告.docx")
    out = _run("doc.to_pdf", source, tmp_path / "out", {"engine": "builtin"}, "报告.pdf")

    assert out.is_file()
    doc = pymupdf.open(str(out))
    try:
        text = "".join(page.get_text() for page in doc)
        assert "季度工作报告" in text
        assert "表格" in text or "文档转换" in text
    finally:
        doc.close()


def test_builtin_engine_handles_text_file(tmp_path: Path) -> None:
    source = tmp_path / "笔记.txt"
    source.write_text("# 会议纪要\n\n讨论了文件转换工具的进度。\n", encoding="utf-8")

    out = _run("doc.to_pdf", source, tmp_path / "out", {"engine": "builtin"}, "笔记.pdf")
    doc = pymupdf.open(str(out))
    try:
        assert "会议纪要" in doc[0].get_text()
    finally:
        doc.close()


def test_builtin_engine_rejects_excel_with_guidance(tmp_path: Path) -> None:
    """Excel/PPT 靠文字提取没有意义，应给出可读的说明而不是产出垃圾。"""
    source = tmp_path / "表.xlsx"
    source.write_bytes(b"placeholder")

    with pytest.raises(ActionError, match="未检测到 Microsoft Office"):
        _run("doc.to_pdf", source, tmp_path / "out", {"engine": "builtin"}, "表.pdf")


def test_auto_falls_back_when_com_unavailable(tmp_path: Path, monkeypatch) -> None:
    """auto 模式下 COM 不可用时应逐级降级，而不是直接失败。"""
    from docforge.actions import doc_convert

    def boom(*_args, **_kwargs):
        from docforge.engines.com import ComUnavailable

        raise ComUnavailable("模拟：Office 不可用")

    monkeypatch.setattr(doc_convert, "_com_word_to_pdf", boom)

    source = _make_docx(tmp_path / "报告.docx")
    out = _run("doc.to_pdf", source, tmp_path / "out", {"engine": "auto"}, "报告.pdf")

    assert out.is_file(), "COM 失败后应降级到内置渲染并成功产出"
    doc = pymupdf.open(str(out))
    try:
        assert "季度工作报告" in doc[0].get_text()
    finally:
        doc.close()


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "archive.zip"
    source.write_bytes(b"PK\x03\x04")
    with pytest.raises(ActionError, match="不支持"):
        _run("doc.to_pdf", source, tmp_path / "out", {}, "out.pdf")


# --------------------------------------------------------------------------- #
# PDF → Word                                                                    #
# --------------------------------------------------------------------------- #

def _make_text_pdf(path: Path, pages: int = 2) -> Path:
    doc = pymupdf.open()
    for index in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 80), f"Annual Report Section {index + 1}", fontsize=18, fontname="helv")
        page.insert_text((60, 120), "This paragraph describes quarterly progress.", fontsize=11, fontname="helv")
        page.insert_text((60, 145), "Revenue grew steadily across all regions.", fontsize=11, fontname="helv")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return path


def test_pdf_to_word_via_pdf2docx(tmp_path: Path) -> None:
    source = _make_text_pdf(tmp_path / "report.pdf", pages=2)
    out = _run("pdf.to_word", source, tmp_path / "out", {}, "report.docx")

    assert out.suffix == ".docx"
    assert out.stat().st_size > 3000

    from docx import Document

    document = Document(str(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Annual Report Section 1" in text
    assert "quarterly progress" in text
    assert "Revenue grew steadily" in text


def test_pdf_to_word_respects_page_range(tmp_path: Path) -> None:
    source = _make_text_pdf(tmp_path / "report.pdf", pages=3)
    out = _run("pdf.to_word", source, tmp_path / "out", {"pages": "2"}, "part.docx")

    from docx import Document

    document = Document(str(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Section 2" in text
    assert "Section 1" not in text
    assert "Section 3" not in text


def test_scanned_pdf_gives_actionable_guidance(tmp_path: Path) -> None:
    """扫描件没有文字层，转换会得到空文档。与其让用户困惑，不如提前说清楚。"""
    from PIL import Image
    import io

    image = Image.new("RGB", (800, 1100), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    source = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=buffer.getvalue())
    doc.save(str(source))
    doc.close()

    with pytest.raises(ActionError, match="扫描件"):
        _run("pdf.to_word", source, tmp_path / "out", {}, "scan.docx")


def test_sparse_text_page_is_not_treated_as_scan(tmp_path: Path) -> None:
    """文字很少但**没有图**的页面不是扫描件。

    这是个真实的误判场景：封面页、标题页往往只有二三十个字。
    如果判据只看"文字少"，用户就会拿到"这看起来是扫描件，请先做 OCR"的提示，
    而他手上的明明是一份正常的文字 PDF。
    """
    source = tmp_path / "cover.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((200, 300), "Annual Report 2026", fontsize=24, fontname="helv")
    doc.save(str(source))
    doc.close()

    out = _run("pdf.to_word", source, tmp_path / "out", {}, "cover.docx")

    from docx import Document

    document = Document(str(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Annual Report 2026" in text


def test_scanned_pdf_can_be_rebuilt_with_ocr(tmp_path: Path) -> None:
    """开启 OCR 后，扫描件也应能重建出 Word 文档。"""
    from PIL import Image, ImageDraw, ImageFont
    import io

    font_path = "C:/Windows/Fonts/msyh.ttc"
    if not Path(font_path).is_file():
        pytest.skip("需要微软雅黑字体来生成测试图")

    image = Image.new("RGB", (1000, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((50, 60), "会议纪要 第三季度", font=ImageFont.truetype(font_path, 40), fill="black")
    draw.text((50, 160), "讨论了文件转换工具的交付进度", font=ImageFont.truetype(font_path, 28), fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    source = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    # 铺满页面并保持比例
    page.insert_image(pymupdf.Rect(40, 40, 555, 246), stream=buffer.getvalue())
    doc.save(str(source))
    doc.close()

    out = _run("pdf.to_word", source, tmp_path / "out", {"ocr_scanned": True}, "scan.docx")

    from docx import Document

    document = Document(str(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "会议纪要" in text, f"OCR 结果应写入文档，实际内容：{text[:200]}"


def test_corrupt_pdf_reports_readable_error(tmp_path: Path) -> None:
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"definitely not a pdf")
    with pytest.raises(ActionError, match="无法打开 PDF"):
        _run("pdf.to_word", bad, tmp_path / "out", {}, "out.docx")


def test_bad_page_range_reports_error(tmp_path: Path) -> None:
    source = _make_text_pdf(tmp_path / "report.pdf", pages=2)
    with pytest.raises(ActionError, match="没有匹配到任何页面"):
        _run("pdf.to_word", source, tmp_path / "out", {"pages": "9-20"}, "out.docx")
